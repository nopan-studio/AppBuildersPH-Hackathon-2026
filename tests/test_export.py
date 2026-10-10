import os
import sys
import unittest
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bookyp import export

try:
    import openpyxl
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None


def sample_result():
    return {
        "file": "receipt1.jpg",
        "merchant": "Cafe",
        "date": "2024-03-05",
        "currency": "USD",
        "items": [
            {"name": "Coffee", "qty": 2, "unit_price": 3.0, "amount": 6.0},
        ],
        "subtotal": 6.0,
        "tax": 0.5,
        "total": 6.5,
        "tax_inclusive": False,
        "ocr_mean_conf": 0.95,
        "status": "accepted",
        "review_reasons": [],
    }


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class ExportTests(unittest.TestCase):
    def test_sheets_and_rows(self):
        data = export.write_xlsx([sample_result()])
        workbook = openpyxl.load_workbook(BytesIO(data))
        self.assertEqual(workbook.sheetnames, ["Receipts", "Items"])

        receipts = workbook["Receipts"]
        self.assertEqual(
            [cell.value for cell in receipts[1]], export.RECEIPT_HEADERS
        )
        self.assertEqual(receipts["A2"].value, "receipt1.jpg")

        items = workbook["Items"]
        self.assertEqual(items["B2"].value, "Coffee")
        self.assertEqual(items["E2"].value, 6.0)

    def test_review_reasons_joined(self):
        result = sample_result()
        result["status"] = "needs_review"
        result["review_reasons"] = ["reason one", "reason two"]
        data = export.write_xlsx([result])
        workbook = openpyxl.load_workbook(BytesIO(data))
        column = export.RECEIPT_HEADERS.index("review_reasons") + 1
        self.assertEqual(
            workbook["Receipts"].cell(row=2, column=column).value,
            "reason one; reason two",
        )

    def test_filename_single_vs_many(self):
        self.assertEqual(export.xlsx_filename([sample_result()]), "receipt1.xlsx")
        self.assertEqual(
            export.xlsx_filename([sample_result(), sample_result()]),
            "bookyp-results.xlsx",
        )


if __name__ == "__main__":
    unittest.main()
