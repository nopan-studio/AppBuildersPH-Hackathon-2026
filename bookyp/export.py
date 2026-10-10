"""Excel (.xlsx) export of scan results.

``write_xlsx`` turns a list of result dicts (the same shape the pipeline and
the API return) into a two-sheet workbook:

    Receipts : one row per receipt (summary fields)
    Items    : one row per line item, linked by file name

``openpyxl`` is imported lazily so the rest of the package stays importable
without it.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Optional, Sequence

RECEIPT_HEADERS = [
    "file",
    "merchant",
    "date",
    "currency",
    "subtotal",
    "tax",
    "total",
    "tax_inclusive",
    "ocr_mean_conf",
    "status",
    "review_reasons",
]

ITEM_HEADERS = ["file", "item", "qty", "unit_price", "amount"]


class ExportError(RuntimeError):
    """Raised when the workbook cannot be produced."""


def _receipt_row(result: Dict[str, Any]) -> List[Any]:
    reasons = result.get("review_reasons") or []
    inclusive = result.get("tax_inclusive")
    return [
        result.get("file"),
        result.get("merchant"),
        result.get("date"),
        result.get("currency"),
        result.get("subtotal"),
        result.get("tax"),
        result.get("total"),
        None if inclusive is None else bool(inclusive),
        result.get("ocr_mean_conf"),
        result.get("status"),
        "; ".join(str(r) for r in reasons) or None,
    ]


def _style_sheet(sheet) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        width = 12
        for cell in sheet[letter]:
            if cell.value is not None:
                width = max(width, min(40, len(str(cell.value)) + 2))
        sheet.column_dimensions[letter].width = width


def write_xlsx(
    results: Sequence[Dict[str, Any]], path: Optional[str] = None
) -> bytes:
    """Build (and optionally save) the workbook; return the raw bytes."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise ExportError(
            "openpyxl is required for Excel export (pip install openpyxl)"
        ) from exc

    workbook = Workbook()

    receipts = workbook.active
    receipts.title = "Receipts"
    receipts.append(RECEIPT_HEADERS)
    for result in results:
        receipts.append(_receipt_row(result))

    items_sheet = workbook.create_sheet("Items")
    items_sheet.append(ITEM_HEADERS)
    for result in results:
        for item in result.get("items") or []:
            items_sheet.append(
                [
                    result.get("file"),
                    item.get("name"),
                    item.get("qty"),
                    item.get("unit_price"),
                    item.get("amount"),
                ]
            )

    _style_sheet(receipts)
    _style_sheet(items_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()

    if path:
        with open(path, "wb") as handle:
            handle.write(data)
    return data


def xlsx_filename(results: Sequence[Dict[str, Any]]) -> str:
    """A friendly download name: single file when one receipt, else 'bookyp-results'."""
    if len(results) == 1:
        base = str(results[0].get("file") or "receipt")
        stem = base.rsplit(".", 1)[0]
        return "{}.xlsx".format(stem)
    return "bookyp-results.xlsx"
